# -*- coding: utf-8 -*-
"""masters.py — 6 种母版布局（ppt-diagram-pipeline skill SKILL.md §Step2）。
每个 layout_* 接收 (prs, spec)，加一张 slide。
spec 通用字段：title / subtitle / footer；按 layout 另有 body/cards/table/image 等。"""

import os
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
import theme
from utils import shapes

DIAG_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "diagrams")


def _bg(slide, color):
    # 用全幅矩形作背景（比 slide.background API 可靠，无版本差异）
    _rect(slide, -Inches(0.05), -Inches(0.05),
          theme.SLIDE_W + Inches(0.1), theme.SLIDE_H + Inches(0.1), color)


def _rect(slide, left, top, w, h, fill, line_color=None, line_w=1):
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, w, h)
    sp.fill.solid()
    sp.fill.fore_color.rgb = theme.rgb(fill) if isinstance(fill, str) else fill
    if line_color:
        sp.line.color.rgb = theme.rgb(line_color)
        sp.line.width = Pt(line_w)
    else:
        sp.line.fill.background()
    sp.shadow.inherit = False
    return sp


def _title_bar(slide, title, subtitle=None):
    _rect(slide, theme.MARGIN_X, theme.MARGIN_Y, Inches(0.14), Inches(0.55), "primary")
    left = theme.MARGIN_X + Inches(0.3)
    tb, tf = shapes.add_textbox(slide, left, theme.MARGIN_Y - Inches(0.05),
                                theme.CONTENT_W - Inches(0.3), Inches(0.7), MSO_ANCHOR.MIDDLE)
    shapes.add_para(tf, title, first=True, size=theme.SIZE_H1, bold=True,
                    color=theme.HEX_COLORS["primary"], space_after=0)
    if subtitle:
        shapes.add_para(tf, subtitle, size=theme.SIZE_SMALL, color=theme.HEX_COLORS["muted"],
                        space_before=2)


def fit_image(slide, path, left, top, max_w, max_h):
    """按原图比例 contain 进 (left,top,max_w,max_h) 矩形并居中。"""
    img = Image.open(path)
    r = img.width / img.height
    aw = max_w
    ah = int(aw / r)
    if ah > max_h:
        ah = max_h
        aw = int(ah * r)
    cl = left + (max_w - aw) // 2
    ct = top + (max_h - ah) // 2
    slide.shapes.add_picture(path, cl, ct, width=aw, height=ah)
    return (aw, ah)


def _footer(slide, text):
    tb, tf = shapes.add_textbox(slide, theme.MARGIN_X, theme.SLIDE_H - Inches(0.4),
                               theme.CONTENT_W, Inches(0.3))
    shapes.add_para(tf, text, first=True, size=9, color=theme.HEX_COLORS["muted"],
                   align=PP_ALIGN.RIGHT)


# ---------- 1. 封面 ----------
def layout_cover(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, "light_bg")
    _rect(slide, 0, 0, theme.SLIDE_W, Inches(2.4), "primary")
    _rect(slide, 0, Inches(2.4), theme.SLIDE_W, Inches(0.08), "accent")
    tb, tf = shapes.add_textbox(slide, theme.MARGIN_X, Inches(0.7), theme.CONTENT_W,
                               Inches(1.6), MSO_ANCHOR.MIDDLE)
    shapes.add_para(tf, spec["title"], first=True, size=40, bold=True,
                   color=theme.HEX_COLORS["white"], align=PP_ALIGN.CENTER)
    if spec.get("subtitle"):
        shapes.add_para(tf, spec["subtitle"], size=theme.SIZE_SUBTITLE,
                       color=theme.HEX_COLORS["white"], align=PP_ALIGN.CENTER, space_before=6)
    tb, tf = shapes.add_textbox(slide, theme.MARGIN_X, Inches(3.4), theme.CONTENT_W,
                               Inches(2.9), MSO_ANCHOR.MIDDLE)
    for i, line in enumerate(spec.get("taglines", [])):
        mult = len(spec.get("taglines", [])) > 2
        shapes.add_para(tf, line, first=(i == 0), size=15 if mult else theme.SIZE_H2,
                       color=theme.HEX_COLORS["dark_text"], align=PP_ALIGN.CENTER, space_after=6)
    tb, tf = shapes.add_textbox(slide, theme.MARGIN_X, Inches(6.7), theme.CONTENT_W, Inches(0.4))
    shapes.add_para(tf, spec.get("date", ""), first=True, size=theme.SIZE_SMALL,
                   color=theme.HEX_COLORS["muted"], align=PP_ALIGN.CENTER)
    return slide


# ---------- 2. 章节过渡 ----------
def layout_section(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, "primary")
    tb, tf = shapes.add_textbox(slide, theme.MARGIN_X, Inches(2.6), theme.CONTENT_W,
                               Inches(2.2), MSO_ANCHOR.MIDDLE)
    shapes.add_para(tf, spec.get("num", ""), first=True, size=64, bold=True,
                   color=theme.HEX_COLORS["accent"], align=PP_ALIGN.LEFT, space_after=4)
    shapes.add_para(tf, spec["title"], size=40, bold=True,
                   color=theme.HEX_COLORS["white"], align=PP_ALIGN.LEFT, space_before=4)
    if spec.get("subtitle"):
        shapes.add_para(tf, spec["subtitle"], size=theme.SIZE_BODY,
                       color=theme.HEX_COLORS["light_bg"], space_before=8)
    _rect(slide, theme.MARGIN_X, Inches(2.55), Inches(1.6), Inches(0.06), "accent")
    return slide


# ---------- 3. 标准内容 ----------
def layout_standard(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, "white")
    _title_bar(slide, spec["title"], spec.get("subtitle"))
    top = theme.MARGIN_Y + Inches(0.75)
    h = theme.SLIDE_H - top - Inches(0.5)
    anchor = MSO_ANCHOR.BOTTOM if spec.get("table") else MSO_ANCHOR.MIDDLE
    tb, tf = shapes.add_textbox(slide, theme.MARGIN_X, top, theme.CONTENT_W, h, anchor)
    for i, p in enumerate(spec.get("body", [])):
        shapes.add_para(tf, p.get("text", ""), first=(i == 0),
                       size=p.get("size", theme.SIZE_BODY),
                       bold=p.get("bold", False), color=p.get("color"),
                       align=p.get("align", PP_ALIGN.LEFT),
                       space_after=p.get("space_after", 7), level=p.get("level", 0))
    if spec.get("table"):
        _add_table(slide, spec["table"], top, theme.CONTENT_W, h)
    if spec.get("footer"):
        _footer(slide, spec["footer"])
    return slide


def _add_table(slide, tspec, top, width, max_h):
    headers = tspec["headers"]
    rows = tspec["rows"]
    ncol = len(headers)
    nrow = len(rows) + 1
    tbl_h = min(max_h, Inches(0.4) * nrow)
    ct = top + (max_h - tbl_h) // 2  # 垂直居中
    gtbl = slide.shapes.add_table(nrow, ncol, theme.MARGIN_X, ct, width, tbl_h)
    tbl = gtbl.table
    # 列宽
    widths = tspec.get("col_widths")
    if widths:
        total = sum(widths)
        for i, w in enumerate(widths):
            tbl.columns[i].width = int(w / total * width)
    # 表头
    for j, htext in enumerate(headers):
        shapes.fill_cell(tbl.cell(0, j), htext, size=theme.SIZE_TABLE, bold=True,
                        color=theme.HEX_COLORS["white"], fill=theme.HEX_COLORS["primary"],
                        align=PP_ALIGN.CENTER)
    # 行
    for i, row in enumerate(rows, 1):
        for j, cell in enumerate(row):
            fill = theme.HEX_COLORS["light_bg"] if i % 2 == 0 else theme.HEX_COLORS["white"]
            shapes.fill_cell(tbl.cell(i, j), cell, size=theme.SIZE_TABLE,
                            fill=fill, align=PP_ALIGN.LEFT)
    return tbl


# ---------- 4. 左文右图 ----------
def layout_text_image(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, "white")
    _title_bar(slide, spec["title"], spec.get("subtitle"))
    top = theme.MARGIN_Y + Inches(0.8)
    h = theme.SLIDE_H - top - Inches(0.5)
    side = spec.get("image_side", "right")
    text_w = Inches(5.6)
    img_w = theme.CONTENT_W - text_w - Inches(0.4)
    if side == "right":
        text_left = theme.MARGIN_X
        img_left = theme.MARGIN_X + text_w + Inches(0.4)
    else:
        text_left = theme.MARGIN_X + img_w + Inches(0.4)
        img_left = theme.MARGIN_X
    tb, tf = shapes.add_textbox(slide, text_left, top, text_w, h, MSO_ANCHOR.MIDDLE)
    for i, p in enumerate(spec.get("body", [])):
        shapes.add_para(tf, p.get("text", ""), first=(i == 0),
                       size=p.get("size", theme.SIZE_BODY), bold=p.get("bold", False),
                       color=p.get("color"), space_after=p.get("space_after", 6),
                       level=p.get("level", 0))
    if spec.get("image"):
        path = os.path.join(DIAG_DIR, spec["image"])
        fit_image(slide, path, img_left, top, img_w, h)
    if spec.get("footer"):
        _footer(slide, spec["footer"])
    return slide


# ---------- 5. 卡片网格 ----------
def layout_cards(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, "white")
    _title_bar(slide, spec["title"], spec.get("subtitle"))
    cards = spec.get("cards", [])
    top = theme.MARGIN_Y + Inches(0.85)
    h = theme.SLIDE_H - top - Inches(0.5)
    cols = spec.get("cols", 3)
    rows_n = (len(cards) + cols - 1) // cols
    gap = Inches(0.25)
    cw = (theme.CONTENT_W - gap * (cols - 1)) / cols
    ch = (h - gap * (rows_n - 1)) / rows_n
    for i, card in enumerate(cards):
        r = i // cols
        c = i % cols
        cl = theme.MARGIN_X + c * (cw + gap)
        ct = top + r * (ch + gap)
        accent = card.get("color", "accent")
        _rect(slide, cl, ct, cw, ch, "light_bg", line_color="line")
        _rect(slide, cl, ct, Inches(0.1), ch, accent)
        tb, tf = shapes.add_textbox(slide, cl + Inches(0.25), ct + Inches(0.12),
                                    cw - Inches(0.4), ch - Inches(0.2))
        shapes.add_para(tf, card["title"], first=True, size=theme.SIZE_BODY, bold=True,
                       color=theme.HEX_COLORS["primary"], space_after=3)
        for j, line in enumerate(card.get("body", [])):
            shapes.add_para(tf, line, size=theme.SIZE_SMALL, color=theme.HEX_COLORS["dark_text"],
                           space_after=2)
    return slide


# ---------- 6. 上下分栏 ----------
def layout_split(prs, spec):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _bg(slide, "white")
    _title_bar(slide, spec["title"], spec.get("subtitle"))
    top = theme.MARGIN_Y + Inches(0.8)
    h = theme.SLIDE_H - top - Inches(0.5)
    img_first = spec.get("img_first", False)
    text_h = Inches(1.6) if spec.get("image") else h
    img_h = h - text_h - Inches(0.2)
    if img_first:
        fit_image(slide, os.path.join(DIAG_DIR, spec["image"]), theme.MARGIN_X, top,
                  theme.CONTENT_W, img_h)
        tb, tf = shapes.add_textbox(slide, theme.MARGIN_X, top + img_h + Inches(0.2),
                                    theme.CONTENT_W, text_h)
    else:
        tb, tf = shapes.add_textbox(slide, theme.MARGIN_X, top, theme.CONTENT_W, text_h)
        if spec.get("image"):
            fit_image(slide, os.path.join(DIAG_DIR, spec["image"]), theme.MARGIN_X,
                      top + text_h + Inches(0.2), theme.CONTENT_W, img_h)
    for i, p in enumerate(spec.get("body", [])):
        shapes.add_para(tf, p.get("text", ""), first=(i == 0),
                       size=p.get("size", theme.SIZE_BODY), bold=p.get("bold", False),
                       color=p.get("color"), space_after=p.get("space_after", 4),
                       level=p.get("level", 0))
    return slide


LAYOUTS = {
    "cover": layout_cover,
    "section": layout_section,
    "standard": layout_standard,
    "text_image": layout_text_image,
    "cards": layout_cards,
    "split": layout_split,
}
