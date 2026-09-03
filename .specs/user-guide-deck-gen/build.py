#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build.py — flow-kit-用户指南.pptx 声明式生成器（2026-09 重建 · user-guide-sync-2026-09）。

用法:
    python3 .specs/user-guide-deck-gen/build.py            # 读 slides.json → 输出根目录 pptx
    python3 .specs/user-guide-deck-gen/build.py --smoke    # 冒烟：3 个占位页输出 /tmp/guide-smoke.pptx

依赖（本机已锁定）:
    python-pptx >= 1.0.2 · Pillow（utils/shapes 不需但保留）· LibreOffice >= 24.2（仅渲染验证用，本脚本不调）
"""
import os
import shutil
import sys
import json
from pathlib import Path

GEN_DIR = Path(__file__).resolve().parent
ROOT = GEN_DIR.parents[1]           # .../flow-kit/.specs/user-guide-deck-gen -> flow-kit 根
SLIDES = GEN_DIR / "slides.json"
OUT = ROOT / "flow-kit-用户指南.pptx"

sys.path.insert(0, str(GEN_DIR))
import theme                          # noqa: E402
from layouts import LAYOUTS           # noqa: E402
from pptx import Presentation         # noqa: E402


def check_env():
    try:
        import pptx as _p
        ver = getattr(_p, "__version__", "unknown")
        assert tuple(int(x) for x in str(ver).split(".")) >= (1, 0, 2), f"python-pptx too old: {ver}"
        print(f"[env] python-pptx {ver} OK")
    except Exception as exc:
        print(f"[env] WARN python-pptx: {exc}")
    if shutil.which("soffice"):
        print("[env] LibreOffice (soffice) found")
    else:
        print("[env] WARN soffice not found — 渲染验证请装 LibreOffice")


def build(slides, out_path):
    prs = Presentation()
    prs.slide_width = theme.SLIDE_W
    prs.slide_height = theme.SLIDE_H
    for i, spec in enumerate(slides, 1):
        layout = spec.get("layout", "band")
        if layout not in LAYOUTS:
            raise SystemExit(f"unknown layout {layout!r} at slide {i}")
        LAYOUTS[layout](prs, spec)
        print(f"  {i:>2} {layout:<6} {spec.get('title','')[:50]}")
    prs.save(out_path)
    reopen = Presentation(out_path)
    print(f"OK: {out_path}  slides={len(reopen.slides)}")


def smoke():
    slides = [
        {"layout": "cover", "title": "Flow-Kit 用户指南（smoke）", "subtitle": "占位副题",
         "taglines": ["tagline 1", "tagline 2"], "date": "2026-09-03 | github.com/hellrabb/flow-kit"},
        {"layout": "band", "title": "版式冒烟", "subtitle": "band layout",
         "bullets": [{"t": "正文层级 1", "size": 16}, {"t": "子要点", "lvl": 1, "c": "accent"},
                     {"t": "强调要点", "bold": True, "c": "danger"}]},
        {"layout": "table", "title": "表格冒烟", "subtitle": "table layout",
         "rows": [["A", "B"], ["1", "2"]], "header": ["列一", "列二"],
         "cols": [0, 1]},
    ]
    out = Path("/tmp/guide-smoke.pptx")
    build(slides, out)
    print("smoke out:", out)


def main():
    check_env()
    if "--smoke" in sys.argv:
        smoke()
        return
    if not SLIDES.exists():
        raise SystemExit(f"missing {SLIDES} —— slides.json 尚未编写（T07 产物）")
    with open(SLIDES, encoding="utf-8") as fh:
        slides = json.load(fh)
    assert isinstance(slides, list) and len(slides) >= 19, "slides.json 应为 >=19 页的 list"
    out = OUT
    if "--out" in sys.argv:
        out = Path(sys.argv[sys.argv.index("--out") + 1])
    build(slides, out)


if __name__ == "__main__":
    main()
