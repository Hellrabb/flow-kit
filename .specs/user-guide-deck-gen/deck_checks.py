#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""deck_checks.py — flow-kit-用户指南.pptx 成品断言（user-guide-sync-2026-09 · T08）。

断言：页数=20；首页含 2026-09-03 与仓库 URL；逐页文本非空（无空页）；
禁词 0（与 FLOW-KIT-用户指南.md AC-2 同清单）；页 14 含五级链字段；
页 20 含 dsh 插件内容；deck 全文含关键串。
"""
import sys
from pathlib import Path
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]            # .../flow-kit/.specs/user-guide-deck-gen -> flow-kit 根
PPTX = ROOT / "flow-kit-用户指南.pptx"
BANNED = ["20260713", "17 个模块", "三级优先级链", "三级链",
          "仅 Claude Code", "仅为 Claude Code", "只为 Claude Code"]
KEY_STRINGS = ["l2-default=", "l3-default=", "dsh plugin", "五级", "archive-commit", "34 号"]
EXPECT_PAGES = 20


def slide_texts():
    prs = Presentation(str(PPTX))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for sh in slide.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text.strip())
            if sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            texts.append(cell.text.strip())
        out.append((i, "\n".join(texts)))
    return out


def main():
    items = slide_texts()
    assert len(items) == EXPECT_PAGES, f"pages={len(items)}"
    full = "\n".join(t for _, t in items)

    import json as _json
    slides_src = _json.dumps(_json.load(open(Path(__file__).resolve().parent / "slides.json", encoding="utf-8")), ensure_ascii=False)
    for b in BANNED:
        assert b not in full, f"banned found: {b}"
        assert b not in slides_src, f"banned in slides.json: {b}"

    page1, p1text = items[0]
    assert "2026-09-03" in p1text, "cover date missing"
    assert "hellrabbit/flow-kit" in p1text, "cover url missing"

    for i, t in items:
        assert t.strip(), f"empty slide: {i}"

    p14 = items[13][1]
    assert "l2-default=" in p14 and "l3-default=" in p14, "page14 five-tier fields missing"
    assert "五级" in p14, "page14 five-tier wording missing"

    p20 = items[19][1]
    assert "dsh plugin" in p20, "page20 dsh install missing"
    assert "/flow doctor" in p20, "page20 doctor self-check missing"

    for k in KEY_STRINGS:
        assert k in full, f"key string missing in deck: {k}"

    print(f"deck_checks OK: {EXPECT_PAGES} pages, banned=0, all pages non-empty, keys present")


if __name__ == "__main__":
    main()
